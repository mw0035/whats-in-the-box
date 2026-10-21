from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import sys, os; sys.path.insert(0, os.path.dirname(__file__)); from content_shim import *

F = "Arial"
HDR = PatternFill("solid", fgColor="2F3B4C")
TINT = {"animal": "FBEFD9", "color": "E3EEF9", "frame": "E8F3E3"}
thin = Side(style="thin", color="D0D0D0")
B = Border(bottom=thin)

def sheets_with(cat, val):
    return ", ".join(f"#{s['id']}" for s in SHEETS if s[cat] == val)

def table(ws, headers, rows, widths, tints=None, start_row=1):
    for j, h in enumerate(headers, 1):
        c = ws.cell(start_row, j, h)
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HDR
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, row in enumerate(rows, start_row + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.font = Font(name=F); c.alignment = Alignment(vertical="top", wrap_text=True); c.border = B
            if tints and tints[i - start_row - 1]:
                c.fill = PatternFill("solid", fgColor=TINT[tints[i - start_row - 1]])
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(start_row + 1, 1)

wb = Workbook()

# ---------------- How it works
ws = wb.active; ws.title = "How it works"
rs = {t: [v for f in FAMILIES if f["track"] == t and f["outcome"][0] == "route" for v in f["starts"]] for t in TRACKS}
ROUTE_TEXT = ("How often routes trigger. Routes start from these values: "
    + "; ".join(f"{TRACK_LABEL[t]} - {', '.join(rs[t])} ({len(rs[t])} of {len(ANIMALS if t=='animal' else COLORS if t=='color' else FRAMES)})" for t in TRACKS)
    + ". Move values between trail families on the Trails tab to make routes rarer or more common.")
lines = [
    ("What's in the box?", "title"),
    ("A game about the division of cognitive labor. Answer: a chocolate coin. Decoys: penny, gold, dollar bill, coin.", None),
    ("", None),
    ("Setup", "h"),
    ("Everyone scans one QR code on the projector. The app asks about color vision (anyone who says yes is never put on the Color track), then assigns a track round-robin: Animal, Color or Frame.", None),
    ("12 printed sheets around the room. Each shows one icon: a colored animal inside a shaped frame, plus a QR code and a sheet number. See the Sheets tab.", None),
    ("Each player's phone has three buttons: Field guide (their own track's vocabulary only), Notebook (clues found so far; clues only, not sheets visited) and Submit my answer. While on a trail, a 'Start a new trail' button lets a player abandon it.", None),
    ("Track wording tells players what their field guide helps with ('Your field guide helps you investigate animals'). Nothing tells them to ignore the rest of the sheet.", None),
    ("", None),
    ("How a trail works", "h"),
    ("1. The player scans a sheet to start a trail. The value of their track on that sheet (its animal, color or frame) picks the trail family. See the Trails tab.", None),
    ("2. Two questions follow. Each is answered by scanning a sheet with the right value. Color and Frame answers accept either of the two sheets with that value.", None),
    ("3. Wrong answer on an ordinary question: 'Your trail goes cold. Return to sheet #N and scan it to try again.' Rescanning sheet #N repeats the same question. (Forgiving.)", None),
    ("4. After two correct answers, most trails end in a clue (a local optimum). The clue goes in the notebook, and the player can start a new trail from any sheet.", None),
    ("5. A few trails end instead in a bare word from another track's field guide ('Next: bloat'). This is a cross-track route. See the Routes tab.", None),
    ("", None),
    ("Cross-track routes (the only way to the answer)", "h"),
    ("Each route has two hops, each using a term from a different other track's field guide. Routes are designed so that complementary expertise makes success much more likely; a player can still finish one alone by knowing the term, looking it up elsewhere, or guessing.", None),
    ("One shot per hop: a wrong scan loses the route and ends the trail ('Your trail goes cold.'). The player starts a new trail from any sheet.", None),
    ("A player can retry a lost route only by starting the whole trail over from its first scan. (Setting in content.json: retryLostRoutes. Set it to false and a lost route stays lost for the rest of the run.)", None),
    ("No timer on hop prompts: the cost falls on guessing, not on taking time to find help.", None),
    ("The final hop reveals: 'You found it. It's a chocolate coin.' Scanning that sheet in any other situation shows ordinary content.", None),
    ("", None),
    ("Clue design", "h"),
    ("Every clue is true of a chocolate coin. Each track's clues are written to suggest its decoy (Animal → penny, Color → gold, Frame → dollar bill), and the Clues tab says what each clue is meant to suggest. The tab checks compatibility only: pooled, the clues rule out penny, gold and dollar bill but leave coin and chocolate coin, so no local clue can separate the two. Whether students actually find each decoy the most convincing answer needs a playtest.", None),
    ("", None),
    ("Submitting answers", "h"),
    ("Players can submit one answer at any time. Submitting ends that player's scanning for the run, but they can still talk with anyone and keep their field guide and notebook, so their expertise stays available to others.", None),
    ("The projector shows a running count: 'Run 1: 12 of 30 have submitted.'", None),
    ("Answers are free text, scored by closeness to the answer (see the Scoring tab). When everyone has submitted, or the instructor presses Reveal scores now, each phone shows that player's own answer and score, and the projector shows the distribution of scores. Nothing reveals the answer itself, so the class can play again.", None),
    ("", None),
    ("Runs", "h"),
    ("The game can be played in several runs (for example, Run 1 in silence and Run 2 with talking allowed). The instructor starts each run from the instructor screen. Talking rules are announced aloud; the app doesn't enforce them.", None),
    ("Between runs, players keep their track and their notebook. Trails and routes reset, so a route lost in Run 1 can be found again in Run 2, and a strange word a player couldn't decode in silence becomes something they can now ask about.", None),
    ("", None),
    ("No hints before a route", "h"),
    ("Nothing a player can see before completing a cross-track route points toward chocolate, candy or sweets: no clue, riddle, fallback, field-guide note or message. The only exception is the final screen of a completed route. check.py enforces this with a list of forbidden words.", None),
    ("", None),
    ("Still to tune (simulation)", "h"),
    (ROUTE_TEXT, None),
    ("", None),
    ("Editing", "h"),
    ("This workbook is generated from the app's content.json, which is what the game actually reads. Edit content.json (or ask Claude), then regenerate this workbook with: python3 tools/build_workbook.py", None),
    ("Every tab is meant to be edited: riddle wording, clue text, guide notes and image prompts. If you change an answer, a term or a route, keep the rules above true: every answer has a sheet, each route term appears in exactly one guide (not the receiver's), and every clue stays true of a chocolate coin.", None),
]
for i, (t, kind) in enumerate(lines, 1):
    c = ws.cell(i, 1, t); c.alignment = Alignment(wrap_text=True, vertical="top")
    c.font = Font(name=F, size=16, bold=True) if kind == "title" else Font(name=F, size=12, bold=True) if kind == "h" else Font(name=F)
ws.column_dimensions["A"].width = 110

# ---------------- Sheets
ws = wb.create_sheet("Sheets")
rows = [[s["id"], s["animal"], s["color"], s["frame"], image_prompt(s)] for s in SHEETS]
table(ws, ["Sheet #", "Animal", "Color", "Frame", "Image prompt"], rows, [9, 14, 11, 18, 80])

# ---------------- Trails
ws = wb.create_sheet("Trails")
rows, tints = [], []
for f in FAMILIES:
    t = f["track"]
    cat = t
    def ans(a):
        if a == START: return "the start value (loop)"
        return f"{a}  ({sheets_with(cat, a)})"
    kind, ref = f["outcome"]
    if kind == "clue":
        out_t, out = "Clue (local optimum)", f"{CLUES[ref]['text']}  [{ref}]"
    else:
        r = next(x for x in ROUTES if x["id"] == ref)
        out_t, out = "Cross-track route", f"{ref}: Next: {r['hops'][0]['term']}"
    rows.append([f["id"], TRACK_LABEL[t], ", ".join(f["starts"]), f["steps"][0][0], ans(f["steps"][0][1]),
                 f["steps"][1][0], ans(f["steps"][1][1]), out_t, out])
    tints.append(t)
table(ws, ["Trail", "Track", "Starts from", "Question 1", "Answer 1 (sheets)", "Question 2", "Answer 2 (sheets)", "Ends in", "Outcome"],
      rows, [7, 9, 22, 34, 22, 34, 22, 18, 34], tints)

# ---------------- Routes
ws = wb.create_sheet("Routes")
rows, tints = [], []
for r in ROUTES:
    for k, h in enumerate(r["hops"], 1):
        cat, val = h["accept"]
        rows.append([r["id"], TRACK_LABEL[r["receiver"]], r["family"], k, f"Next: {h['term']}", TRACK_LABEL[h["owner"]],
                     f"{val}  ({sheets_with(cat, val)})", ANSWER if k == 2 else "Next hop"])
        tints.append(r["receiver"])
table(ws, ["Route", "Player's track", "From trail", "Hop", "Shown on phone", "Decoded with", "Correct scan (sheets)", "If right"],
      rows, [7, 10, 8, 5, 16, 11, 24, 22], tints)

# ---------------- Clues
ws = wb.create_sheet("Clues")
rows, tints = [], []
for k, c in CLUES.items():
    rows.append([k, TRACK_LABEL[c["track"]], c["text"], SUGGESTS[k]] + ["✓" if c["fits"][x] else "✗" for x in CANDIDATES])
    tints.append(c["track"])
table(ws, ["ID", "Track", "Clue", "Written to suggest", "Fits penny?", "Fits gold?", "Fits dollar bill?", "Fits coin?", "Fits chocolate coin?"], rows,
      [10, 9, 44, 20, 11, 10, 15, 10, 18], tints)
n = len(rows); last = n + 1; r0 = last + 2
ws.cell(r0, 4, "Fits every clue?").font = Font(name=F, bold=True)
ws.cell(r0 + 1, 4, "Clues it fails").font = Font(name=F, bold=True)
for j in range(5, 10):
    L = get_column_letter(j)
    ws.cell(r0, j, f'=IF(COUNTIF({L}2:{L}{last},"✗")=0,"yes","no")').font = Font(name=F, bold=True)
    ws.cell(r0 + 1, j, f'=COUNTIF({L}2:{L}{last},"✗")').font = Font(name=F)
ws.cell(r0 + 3, 1, "These columns check compatibility, not persuasiveness: a decoy can fit every clue from its track and still lose to another answer in students' minds. Pooled, only coin and chocolate coin survive; only a cross-track route separates them, and no clue mentions or suggests anything edible.").font = Font(name=F, italic=True)

# ---------------- Guides
for t, title in [("animal", "Guide - Animal"), ("color", "Guide - Color"), ("frame", "Guide - Frame")]:
    ws = wb.create_sheet(title)
    route_terms = {h["term"].lower(): r["id"] for r in ROUTES for h in r["hops"] if h["owner"] == t}
    rows = []
    for term, val, note in GUIDES[t]:
        key = next((k for k in route_terms if term.lower() == k or term.lower().startswith(k + " ")), None)
        rows.append([term, note, val or "—", sheets_with(t, val) if val else "— (padding)",
                     f"Yes ({route_terms[key]})" if key else ""])
    table(ws, ["Term", "Meaning (shown in app)", "Printed value", "Sheets", "Route term?"], rows, [20, 56, 16, 18, 13])
    if t == "color":
        r0 = len(rows) + 3
        ws.cell(r0, 1, "Mixing chart (also in the Color field guide)").font = Font(name=F, bold=True)
        for i, m in enumerate(COLOR_MIXING, 1):
            ws.cell(r0 + i, 1, m).font = Font(name=F)

# ---------------- Scoring
ws = wb.create_sheet("Scoring")
rows = [[sc, label, rule] for sc, label, rule, ex in SCORING]
table(ws, ["Points", "Category", "Matching rule (checked top to bottom; one typo allowed in longer words)"], rows, [8, 34, 100])
ws.cell(len(rows) + 4, 1, "This tab is generated from content.json. Edit points and words there; anything unmatched appears on the admin page before the reveal.").font = Font(name=F, italic=True)

# ---------------- Messages
ws = wb.create_sheet("Messages")
table(ws, ["Where", "Text on screen"], [list(m) for m in MESSAGES], [30, 100])

wb.save(sys.argv[1] if len(sys.argv) > 1 else "whats_in_the_box_content.xlsx")
print("saved")
